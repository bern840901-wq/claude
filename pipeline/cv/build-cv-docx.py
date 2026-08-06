# -*- coding: utf-8 -*-
# Print twin of twkrpuente.web.app/cv/ — content parsed verbatim from the
# page's embedded DATA (all four languages, full 72-case timeline).
# Style maps the page's CSS 1:1 (px → pt × 0.75): white ground, hairline
# section rules, Cormorant-family serif (EB Garamond) + Inter + Noto CJK.
import json, os, re
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SP = os.path.dirname(os.path.abspath(__file__))
OUT = f'{SP}/cv-docx'
os.makedirs(OUT, exist_ok=True)

page = open('/home/user/claude/cv/index.html', encoding='utf-8').read()
m = re.search(r'const DATA = (\{.*?\});\n', page, re.S)
DATA = json.loads(m.group(1))

INK, MID, LIGHT, HAIR, AC = '14110D', '555555', '8F8A80', 'E9E5DC', 'E8410E'
SERIF, SANS = 'EB Garamond', 'Inter'
CJK_SERIF = {'zh': 'Noto Serif CJK TC', 'ko': 'Noto Serif CJK KR', 'en': 'Noto Serif CJK TC', 'es': 'Noto Serif CJK TC'}
CJK_SANS = {'zh': 'Noto Sans CJK TC', 'ko': 'Noto Sans CJK KR', 'en': 'Noto Sans CJK TC', 'es': 'Noto Sans CJK TC'}
FNAME = {'zh': '游宏斌_CV_精選版', 'ko': '游宏斌_CV_KO', 'en': '游宏斌_CV_EN', 'es': '游宏斌_CV_ES'}
strip = lambda s: re.sub(r'<[^>]+>', '', s or '')
track = lambda em, size: int(em * size * 20)   # letter-spacing em → w:spacing (1/20 pt)

# QR — ink on white, quiet, lives in the footer (the page's foot links → print URL)
import qrcode
from PIL import Image
q = qrcode.QRCode(error_correction=qrcode.constants.ERROR_CORRECT_M, box_size=10, border=2)
q.add_data('https://twkrpuente.web.app/'); q.make(fit=True)
q.make_image(fill_color=(20, 17, 13), back_color=(255, 255, 255)).convert('RGB') \
 .resize((420, 420), Image.NEAREST).save(f'{OUT}/qr.png')


def set_font(run, size, color=MID, bold=False, name=SANS, east=None, sp=None, caps=False, ul=None):
    run.font.size = Pt(size); run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.name = name
    r = run._element.get_or_add_rPr()
    rf = r.find(qn('w:rFonts'))
    if rf is None: rf = OxmlElement('w:rFonts'); r.append(rf)
    for a in ('w:ascii', 'w:hAnsi'): rf.set(qn(a), name)
    for a in ('w:eastAsia', 'w:cs'): rf.set(qn(a), east or name)
    if sp:
        e = OxmlElement('w:spacing'); e.set(qn('w:val'), str(sp)); r.append(e)
    if caps:
        r.append(OxmlElement('w:caps'))
    if ul:
        e = OxmlElement('w:u'); e.set(qn('w:val'), 'single'); e.set(qn('w:color'), ul); r.append(e)


def para(doc, before=0, after=0, line=None, align=None, indent=None):
    p = doc.add_paragraph()
    f = p.paragraph_format
    f.space_before = Pt(before); f.space_after = Pt(after)
    if line: f.line_spacing = line
    if align: p.alignment = align
    if indent: f.left_indent = Cm(indent)
    return p


def hairline_top(p, space='14'):
    pPr = p._p.get_or_add_pPr()
    bd = OxmlElement('w:pBdr'); e = OxmlElement('w:top')
    e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), HAIR)
    e.set(qn('w:space'), space)
    bd.append(e); pPr.append(bd)


def slabel(doc, text, lang, before=21):
    # .slabel: 11px uppercase .22em light, hairline above, ~14px padding-top
    p = para(doc, before=before, after=8)
    hairline_top(p)
    r = p.add_run(text)
    set_font(r, 8.25, LIGHT, False, SANS, CJK_SANS[lang], sp=track(.22, 8.25), caps=True)
    return p


def no_borders(tbl):
    tblPr = tbl._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for tag in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{tag}'); e.set(qn('w:val'), 'none'); b.append(e)
    tblPr.append(b)


def fixed_widths(tbl, widths):
    tbl.autofit = False
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = tbl._tbl.tblPr
    layout = OxmlElement('w:tblLayout'); layout.set(qn('w:type'), 'fixed'); tblPr.append(layout)
    grid = tbl._tbl.find(qn('w:tblGrid'))
    for gc, w in zip(grid.findall(qn('w:gridCol')), widths): gc.set(qn('w:w'), str(w))
    for row in tbl.rows:
        for c, w in zip(row.cells, widths): c.width = Emu(w * 635)


def tcmar(cell, top, bottom, left=0, right=0):
    tcPr = cell._tc.get_or_add_tcPr()
    mm = OxmlElement('w:tcMar')
    for tag, v in (('top', top), ('bottom', bottom), ('left', left), ('right', right)):
        e = OxmlElement(f'w:{tag}'); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); mm.append(e)
    tcPr.append(mm)


def cantsplit(tbl):
    for tr in tbl.rows:
        trPr = tr._tr.get_or_add_trPr()
        trPr.append(OxmlElement('w:cantSplit'))


def rowtable(doc, entries, lang, widths=(1050, 6060, 2130)):
    # the page's .row: [yr | ttl | typ] — 7px vertical padding, no rules
    KS, KA = CJK_SERIF[lang], CJK_SANS[lang]
    tb = doc.add_table(rows=len(entries), cols=3); no_borders(tb)
    fixed_widths(tb, list(widths)); cantsplit(tb)
    for i, (yr, ttl, typ) in enumerate(entries):
        cells = tb.rows[i].cells
        for c in cells: tcmar(c, 55, 55, 0, 60)
        r = cells[0].paragraphs[0].add_run(yr); set_font(r, 9.4, LIGHT, False, SANS, KA)
        p = cells[1].paragraphs[0]; p.paragraph_format.line_spacing = 1.24
        r = p.add_run(ttl); set_font(r, 11.5, INK, True, SERIF, KS)
        p3 = cells[2].paragraphs[0]; p3.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p3.add_run(typ); set_font(r, 7.5, LIGHT, False, SANS, KA, sp=track(.14, 7.5), caps=True)
    return tb


def sub_note(doc, text, lang, indent=1.64, before=0, after=3):
    p = para(doc, before=before, after=after, indent=indent, line=1.35)
    r = p.add_run(text); set_font(r, 9.4, LIGHT, False, SANS, CJK_SANS[lang])
    return p



def add_link(p, text, url, size, color, name, east, ulcolor='C9C4BA'):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.text.run import Run
    r_id = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hl = OxmlElement('w:hyperlink'); hl.set(qn('r:id'), r_id)
    wr = OxmlElement('w:r'); wr.append(OxmlElement('w:rPr'))
    t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = text; wr.append(t)
    hl.append(wr); p._p.append(hl)
    run = Run(wr, p)
    set_font(run, size, color, False, name, east, ul=ulcolor)
    return run


def build(lang):
    d = DATA[lang]
    KS, KA = CJK_SERIF[lang], CJK_SANS[lang]
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = Cm(1.9); sec.bottom_margin = Cm(2.4)
    sec.left_margin = sec.right_margin = Cm(2.67)
    doc.styles['Normal'].font.name = SANS
    doc.styles['Normal'].paragraph_format.space_after = Pt(0)

    # ── header ──
    p = para(doc, after=4)
    r = p.add_run(d['name']); set_font(r, 33, INK, True, SERIF, KS)
    p = para(doc, after=10)
    r = p.add_run(d['nameSub']); set_font(r, 9.75, LIGHT, False, SANS, KA, sp=track(.14, 9.75))
    p = para(doc, after=7, line=1.7)
    r = p.add_run(strip(d['role'])); set_font(r, 11.25, MID, False, SANS, KA)
    p = para(doc, after=14)
    r = p.add_run(d['tagline']); set_font(r, 9, LIGHT, False, SANS, KA, sp=track(.16, 9), caps=True)
    p = para(doc)
    r = p.add_run(d['based'] + ' · ' + d['avail']); set_font(r, 10.1, MID, False, SANS, KA)

    # ── contact ──
    slabel(doc, d['ui']['contact'], lang, before=26)
    p = para(doc, after=5)
    r = p.add_run('twkrbridge@gmail.com'); set_font(r, 15.75, INK, True, SERIF, KS, ul=AC)
    p = para(doc)
    r = p.add_run(strip(d['reply'])); set_font(r, 9.4, LIGHT, False, SANS, KA)

    # ── long-term partners ──
    slabel(doc, d['ui']['partners'], lang)
    for pr in d['partners']:
        rowtable(doc, [(pr['period'].split('–')[0], pr['org'], pr['period'])], lang)
        sub_note(doc, pr['desc'], lang)

    # ── work (full timeline, year shown once per group) ──
    slabel(doc, d['ui']['work'], lang)
    p = para(doc, after=9, line=1.4)
    r = p.add_run(strip(d['tlSub'])); set_font(r, 9.4, LIGHT, False, SANS, KA)
    entries, last = [], None
    for row in d['tl']:
        y = '' if row['y'] == last else str(row['y'])
        last = row['y']
        entries.append((y, row['title'], d['types'].get(row['t'], '')))
    rowtable(doc, entries, lang)

    # ── clients ──
    slabel(doc, d['ui']['clients'], lang)
    for grp in d['clients']:
        p = para(doc, after=6)
        r = p.add_run(grp['title']); set_font(r, 9, MID, True, SANS, KA, sp=track(.1, 9))
        for it in grp['items']:
            p = para(doc, after=1.5, line=1.2)
            r = p.add_run(strip(it['name'])); set_font(r, 10.5, INK, False, SANS, KA)
            if it.get('sub'):
                r = p.add_run('   ' + strip(it['sub'])); set_font(r, 9, LIGHT, False, SANS, KA)
        para(doc, after=7)

    # ── education ──
    slabel(doc, d['ui']['education'], lang)
    for e in d['edu']:
        rowtable(doc, [(str(e['year']).split('–')[0], e['school'], '')], lang)
        sub_note(doc, strip(e['deg']), lang)

    # ── languages ──
    slabel(doc, d['ui']['languages'], lang)
    for x in d['langs']:
        p = para(doc, before=3, after=3, line=1.45)
        r = p.add_run(strip(x['name'])); set_font(r, 9.75, LIGHT, True, SANS, KA, sp=track(.14, 9.75))
        r = p.add_run(' · '); set_font(r, 9.75, LIGHT, False, SANS, KA)
        r = p.add_run(strip(x['level'])); set_font(r, 9.4, LIGHT, False, SANS, KA)

    # ── certs / awards ──
    slabel(doc, d['ui']['certs'], lang)
    for c in d['certs']:
        p = para(doc, before=3, after=1)
        r = p.add_run(c['name']); set_font(r, 9.75, LIGHT, True, SANS, KA, sp=track(.14, 9.75))
        p = para(doc, after=4)
        r = p.add_run(strip(c['issuer'])); set_font(r, 9.4, LIGHT, False, SANS, KA)

    # ── foot: the page's footer links become the printed URL + QR ──
    ft = doc.add_table(rows=1, cols=2); no_borders(ft)
    fixed_widths(ft, [7440, 1800])
    c0, c1 = ft.rows[0].cells
    tcmar(c0, 170, 0, 0, 0); tcmar(c1, 110, 0, 0, 0)
    p = c0.paragraphs[0]; hairline_top(p, space='10')
    add_link(p, d['cta2'], 'https://twkrpuente.web.app/home/', 9.4, LIGHT, SANS, KA)
    r = p.add_run('   ·   '); set_font(r, 9.4, LIGHT, False, SANS, KA)
    add_link(p, '▶ ' + d['ui']['film'], 'https://twkrpuente.web.app/', 9.4, LIGHT, SANS, KA)
    p2 = c0.add_paragraph(); p2.paragraph_format.space_before = Pt(4)
    r = p2.add_run('twkrpuente.web.app   ·   © 游宏斌 YU HUNG PIN')
    set_font(r, 9.4, LIGHT, False, SANS, KA)
    pq = c1.paragraphs[0]; pq.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pq.add_run().add_picture(f'{OUT}/qr.png', width=Cm(1.4))

    doc.save(f'{OUT}/{FNAME[lang]}.docx')
    return f'{OUT}/{FNAME[lang]}.docx'


for lg in ['zh', 'ko', 'en', 'es']:
    print('built', build(lg))
